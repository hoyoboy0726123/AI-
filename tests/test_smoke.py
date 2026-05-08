"""Pre-flight smoke + integration tests.

跑法：python tests/test_smoke.py

涵蓋：
- 模組 import 全通
- 設定檔持久化 / 損毀降級
- 工具註冊（25 個）+ schema JSON-Schema 合法
- ReportGenerator 全流程（含圖片欄位）
- 範本編輯（rename / insert / image）
- 預算 / 對話流程（with stub LLM）
- VLM reviewer 流程
- 圖片資料夾配對
- LLMClient 缺 API key / endpoint 不可達時優雅退化

不涵蓋（需 Windows 與真實 LLM）：
- Word COM
- 真實 Tk 視窗渲染
- 真實 Gemini / Ollama 呼叫
- DnD 拖拉
- 全域快捷鍵
"""

import json
import os
import shutil
import sys
import tempfile
import threading
import types
from pathlib import Path

# ============================================================
# Stub Tk / Windows-only modules so app.* can be imported on Linux
# ============================================================

def install_stubs():
    tk_pkg = types.ModuleType("tkinter")
    tk_pkg.__path__ = []
    sys.modules.setdefault("tkinter", tk_pkg)

    fd_stub = types.ModuleType("tkinter.filedialog")
    for n in ("askopenfilename", "askdirectory", "askopenfilenames", "asksaveasfilename"):
        setattr(fd_stub, n, lambda **kw: "")
    sys.modules["tkinter.filedialog"] = fd_stub
    tk_pkg.filedialog = fd_stub

    mb_stub = types.ModuleType("tkinter.messagebox")
    mb_stub.showinfo = lambda *a, **kw: None
    mb_stub.showwarning = lambda *a, **kw: None
    sys.modules["tkinter.messagebox"] = mb_stub
    tk_pkg.messagebox = mb_stub

    class _W:
        def __init__(self, *a, **kw): pass
        def __getattr__(self, _): return lambda *a, **kw: None

    class _CtkInputDialog(_W):
        def get_input(self): return None

    class _StrVar:
        def __init__(self, value=""): self._v = value
        def get(self): return self._v
        def set(self, v): self._v = v

    class _BoolVar:
        def __init__(self, value=False): self._v = bool(value)
        def get(self): return self._v
        def set(self, v): self._v = bool(v)

    ctk = types.ModuleType("customtkinter")
    for n in ("CTk CTkToplevel CTkLabel CTkButton CTkEntry CTkFrame CTkTextbox "
             "CTkComboBox CTkOptionMenu CTkCheckBox CTkRadioButton CTkProgressBar "
             "CTkScrollableFrame CTkTabview").split():
        setattr(ctk, n, _W)
    ctk.CTkInputDialog = _CtkInputDialog
    ctk.StringVar = _StrVar
    ctk.BooleanVar = _BoolVar
    ctk.IntVar = _StrVar
    ctk.CTkFont = lambda **kw: object()
    ctk.set_appearance_mode = lambda *a: None
    ctk.set_default_color_theme = lambda *a: None
    sys.modules["customtkinter"] = ctk

    for s in ("keyboard", "win32com", "win32com.client"):
        if s not in sys.modules:
            m = types.ModuleType(s)
            if s == "win32com":
                m.__path__ = []
            sys.modules[s] = m
    sys.modules["keyboard"].add_hotkey = lambda *a, **kw: None
    sys.modules["keyboard"].remove_hotkey = lambda *a, **kw: None
    sys.modules["win32com.client"].GetActiveObject = lambda x: None
    sys.modules["win32com.client"].DispatchEx = lambda x: None
    sys.modules["win32com"].client = sys.modules["win32com.client"]


install_stubs()
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


# ============================================================
# Real Tk-free fixtures
# ============================================================

def make_fixtures():
    import pandas as pd
    from docx import Document
    from PIL import Image

    tmp = tempfile.mkdtemp(prefix="autoreport_test_")

    img_dir = os.path.join(tmp, "photos")
    os.makedirs(img_dir)
    images = {}
    for name in ["流程圖.png", "客戶照片_張三.png"]:
        p = os.path.join(img_dir, name)
        Image.new("RGB", (50, 50), color=(0, 100, 200)).save(p)
        images[name] = p

    xlsx = os.path.join(tmp, "data.xlsx")
    df1 = pd.DataFrame({
        "客戶名稱": ["張三", "李四", "王五"],
        "金額": [10000, 25000, 5000],
        "日期": ["2026-01-01", "2026-02-15", "2026-03-30"],
        "照片": [images["客戶照片_張三.png"], "", ""],
    })
    df2 = pd.DataFrame({"備用": [1, 2]})
    with pd.ExcelWriter(xlsx) as w:
        df1.to_excel(w, sheet_name="客戶資料", index=False)
        df2.to_excel(w, sheet_name="其他", index=False)

    docx = os.path.join(tmp, "template.docx")
    doc = Document()
    doc.add_heading("客戶報告", 0)
    doc.add_paragraph("客戶姓名：{{ 客戶名稱 }}")
    doc.add_paragraph("金額：{{ 金額 }} 元")
    doc.add_paragraph("日期：{{ 日期 }}")
    doc.add_paragraph("照片：{{ 照片 }}")
    doc.add_paragraph("備註：（請補充）")
    doc.add_paragraph("圖 1：流程示意")
    doc.save(docx)

    return tmp, docx, xlsx, img_dir, images


# ============================================================
# FakeApp for AppContext
# ============================================================

class FakeVar:
    def __init__(self, v=""): self._v = v
    def get(self): return self._v
    def set(self, v): self._v = v

class FakeBoolVar(FakeVar):
    def __init__(self, v=False): self._v = bool(v)

class _FakeOrch:
    def __init__(self): self._cancel = threading.Event()

class _FakeRubric:
    def get(self, *a): return ""

class _FakeWidget:
    def set(self, v): pass
    def configure(self, **kw): pass


def make_fake_app():
    from app.agent.budget import BudgetTracker

    class FakeApp:
        def __init__(self):
            for n in ["word_path","excel_path","sheet_name","output_dir","filename_template",
                     "llm_provider","gemini_planner_model","gemini_reviewer_model",
                     "ollama_endpoint","ollama_planner_model","ollama_reviewer_model"]:
                setattr(self, n, FakeVar(""))
            self.header_row = FakeVar("1")
            self.image_width_mm = FakeVar("80")
            self.grid_columns = FakeVar("2")
            self.enable_review = FakeBoolVar(False)
            self.review_sampling = FakeVar("100")
            self.max_review_retries = FakeVar("3")
            self.is_generating = False
            self.cancel_event = threading.Event()
            self.agent_orchestrator = _FakeOrch()
            self.rubric_box = _FakeRubric()
            self.budget = BudgetTracker(planner_limit=10, reviewer_limit=10)
            self.btn_generate = _FakeWidget()
            self.progress = _FakeWidget()
            self.progress_label = _FakeWidget()
        def after(self, _d, fn): fn()
        def _header_row_int(self): return 1
        def _image_width_mm_int(self): return 80
        def _grid_columns_int(self): return 2
        def _review_sampling_int(self): return 100
        def _refresh_sheet_list(self, silent=False): pass
        def _agent_set_status(self, s): pass
        def _reset_generation_ui(self): pass

    return FakeApp()


# ============================================================
# Sections
# ============================================================

def section(title):
    print(f"\n[{title}]")


def test_imports():
    section("imports")
    import app.config, app.settings, app.filename, app.generator, app.mapper, app.hotkey
    import app.agent.budget, app.agent.context, app.agent.dialogs, app.agent.docx_render
    import app.agent.folder_scan, app.agent.mapping_suggester, app.agent.orchestrator
    import app.agent.registry, app.agent.reviewer, app.agent.template_edit, app.agent.tools
    import app.agent.llm.base, app.agent.llm.gemini, app.agent.llm.ollama
    import app.ui
    import main
    print("  21 app modules + main loaded ok")


def test_settings():
    section("settings")
    import app.settings as s
    tmp = tempfile.mkdtemp()
    cases = {
        "missing": (Path(tmp) / "missing.json", None),
        "corrupt": (Path(tmp) / "corrupt.json", "{not json"),
        "empty": (Path(tmp) / "empty.json", ""),
        "partial": (Path(tmp) / "partial.json", '{"word_path":"/x.docx"}'),
    }
    for name, (path, content) in cases.items():
        if content is not None:
            path.write_text(content, encoding="utf-8")
        s.SETTINGS_PATH = path
        d = s.load_settings()
        assert d["llm_provider"] == "Gemini"
        assert d["max_planner_calls"] == 50
        assert d["appearance_mode"] == "System"
    print(f"  fallback to defaults on missing/corrupt/empty/partial: ok")

    # Round-trip with Unicode
    p = Path(tmp) / "rt.json"
    s.SETTINGS_PATH = p
    d = s.load_settings()
    d["filename_template"] = "{客戶}_{日期}.docx"
    d["review_rubric"] = "中文 rubric 內容"
    s.save_settings(d)
    re_loaded = s.load_settings()
    assert re_loaded["filename_template"] == "{客戶}_{日期}.docx"
    assert re_loaded["review_rubric"] == "中文 rubric 內容"
    print(f"  Unicode round-trip: ok")


def test_filename():
    section("filename rendering")
    from app.filename import render_filename
    cases = [
        ("{客戶}_{index}.docx", {"客戶": "張三"}, 1, "張三_1.docx"),
        ("{客戶}.docx", {"客戶": "a/b\\c"}, 1, "a_b_c.docx"),
        ("{缺}.docx", {"有": "x"}, 5, "report_5.docx"),
        ("plain", {}, 1, "plain.docx"),
        ("{val}_{index}.docx", {"val": None}, 3, "_3.docx"),
        ("", {}, 7, "report_7.docx"),
    ]
    for tmpl, row, idx, exp in cases:
        got = render_filename(tmpl, row, idx)
        assert got == exp, f"{tmpl} {row} → {got!r} != {exp!r}"
    print(f"  {len(cases)} cases ok")


def test_generator(docx, xlsx):
    section("ReportGenerator")
    from app.generator import ReportGenerator
    out = tempfile.mkdtemp(prefix="gen_")
    gen = ReportGenerator(
        word_path=docx, excel_path=xlsx, output_dir=out,
        sheet_name="客戶資料", header_row=1,
        filename_template="{客戶名稱}.docx", image_width_mm=40,
    )
    assert gen.list_sheets() == ["客戶資料", "其他"]
    assert gen.template_variables() == {"客戶名稱", "金額", "日期", "照片"}
    missing, extra = gen.validate()
    assert missing == set() and extra == set()
    p, t = gen.generate()
    assert p == 3 and t == 3
    files = sorted(os.listdir(out))
    assert files == ["張三.docx", "李四.docx", "王五.docx"]

    # First doc has inline image (照片 col was a real path)
    from docx import Document
    d1 = Document(os.path.join(out, "張三.docx"))
    assert len(d1.inline_shapes) == 1
    print(f"  3 docs produced; 張三 has inline image")

    # Cancel test
    e = threading.Event(); e.set()
    p, t = gen.generate(cancel_event=e)
    assert p == 0
    print(f"  cancel before iter respected")

    # Validate detects missing
    from docx import Document as Doc
    bad = os.path.join(out, "bad.docx")
    d = Doc(); d.add_paragraph("{{ 不存在 }}"); d.save(bad)
    gen2 = ReportGenerator(word_path=bad, excel_path=xlsx, output_dir=out, sheet_name="客戶資料")
    m, _ = gen2.validate()
    assert "不存在" in m
    print(f"  validate detects missing var")


def test_registry(docx, xlsx):
    section("tool registry")
    from app.agent.tools import build_default_registry
    from app.agent.context import AppContext

    app = make_fake_app()
    ctx = AppContext(app)

    reg_ro = build_default_registry()
    assert len(reg_ro.schemas()) == 3

    reg = build_default_registry(ctx)
    schemas = reg.schemas()
    assert len(schemas) == 25
    json.dumps(schemas, ensure_ascii=False)
    for s in schemas:
        assert s["parameters"]["type"] == "object"
        assert "properties" in s["parameters"]
    print(f"  read-only=3, full=25, all schemas valid")

    # Each func can be called with empty args without raising
    for tool in reg.all():
        r = reg.run(tool.name, {})
        assert isinstance(r, dict), f"{tool.name}: {type(r)}"
    print(f"  all 25 tools handle empty args without raising")


def test_template_edit(docx, images):
    section("template editing")
    from app.agent.template_edit import (
        read_docx_text, rename_template_variable,
        insert_template_variable, insert_image_at_anchor,
    )
    from docx import Document

    work = tempfile.mktemp(suffix=".docx")
    shutil.copy(docx, work)

    r = read_docx_text(work)
    assert any("{{ 客戶名稱 }}" in p for p in r["paragraphs"])

    r = rename_template_variable(work, "客戶名稱", "Customer")
    assert r["changed"] >= 1

    text = "\n".join(read_docx_text(work)["paragraphs"])
    assert "{{ Customer }}" in text and "{{ 客戶名稱 }}" not in text
    print(f"  rename: ok")

    r = insert_template_variable(work, "備註：", "備註內容", "after")
    assert r["inserted"]
    assert any("備註：{{ 備註內容 }}" in p for p in read_docx_text(work)["paragraphs"])
    print(f"  insert: ok")

    img = list(images.values())[0]
    r = insert_image_at_anchor(work, "圖 1：流程示意", img, width_mm=50)
    assert r["inserted"]
    d = Document(work)
    assert len(d.inline_shapes) == 1
    print(f"  insert_image_at_anchor: ok ({len(d.inline_shapes)} shape)")


def test_folder_scan():
    section("folder scan")
    from app.agent.folder_scan import list_folder_files
    tmp = tempfile.mkdtemp()
    for n in ["a.png", "b.JPG", "c.gif", "d.txt"]:
        open(os.path.join(tmp, n), "w").close()
    os.makedirs(os.path.join(tmp, "subdir"))

    r = list_folder_files(tmp, kind="image")
    assert sorted(f["name"] for f in r["files"]) == ["a.png", "b.JPG", "c.gif"]

    r = list_folder_files(tmp, kind="any", max_files=2)
    assert len(r["files"]) == 2 and r["count"] == 4

    assert "error" in list_folder_files("")
    assert "error" in list_folder_files(tmp, kind="bogus")
    print(f"  case-insensitive, max_files, errors: ok")


def test_reviewer():
    section("reviewer JSON parsing")
    from app.agent.reviewer import _extract_json, move_to_failed_reports
    cases = [
        ('{"a":1}', {"a": 1}),
        ('  ```json\n{"x": 2}\n```  ', {"x": 2}),
        ('preamble {"y": 3} trailing', {"y": 3}),
        ('not json', None),
        ('', None),
    ]
    for text, exp in cases:
        got = _extract_json(text)
        if exp is None:
            assert got is None
        else:
            assert got == exp
    print(f"  {len(cases)} JSON cases ok")

    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "x.docx")
    open(src, "w").close()
    p1 = move_to_failed_reports(src, os.path.join(tmp, "fail"))
    p2 = move_to_failed_reports(src, os.path.join(tmp, "fail"))
    assert os.path.basename(p1) == "x.docx"
    assert "_1" in os.path.basename(p2)
    print(f"  dedup: ok")


def test_suggest_mapping(docx, xlsx, img_dir):
    section("suggest_mappings + suggest_image_placements")
    from app.agent.context import AppContext
    from app.agent.llm.base import Message

    class StubLLM:
        next_text = ""
        def is_available(self): return True
        def chat(self, messages, model=None, tools=None):
            return Message(role="assistant", text=self.next_text)

    app = make_fake_app()
    app.word_path.set(docx)
    app.excel_path.set(xlsx)
    app.sheet_name.set("客戶資料")
    ctx = AppContext(app)
    stub = StubLLM()
    ctx._build_planner_client = lambda: (stub, "stub", None)

    stub.next_text = '{"renames":[{"from":"客戶名稱","to":"Customer"}],"inserts":[]}'
    r = ctx.suggest_mappings()
    assert r["renames"][0]["from"] == "客戶名稱"
    print(f"  suggest_mappings: ok")

    stub.next_text = '{"placements":[{"image":"流程圖.png","anchor":"圖 1：流程示意"}]}'
    r = ctx.suggest_image_placements(image_folder=img_dir)
    assert r["placements"][0]["image"] == "流程圖.png"
    assert r["placements"][0]["image_path"].endswith("流程圖.png")
    print(f"  suggest_image_placements: ok")


def test_review_flow(docx, xlsx):
    section("review flow + budget exhaustion")
    import app.agent.reviewer as reviewer_mod
    from PIL import Image as _Image
    from app.agent.context import AppContext
    from app.agent.budget import BudgetTracker

    # Fake renderer that uses its own tempdir (so cleanup doesn't wipe fixtures)
    def fake_render(*a, **kw):
        d = tempfile.mkdtemp(prefix="fake_render_")
        img = os.path.join(d, "page_001.png")
        _Image.new("RGB", (50, 50), color="white").save(img)
        return [(1, img)]
    reviewer_mod.docx_to_images = fake_render

    class StubVLM:
        responses = []
        idx = 0
        def is_available(self): return True
        def vision_complete(self, **kw):
            txt = StubVLM.responses[StubVLM.idx % len(StubVLM.responses)]
            StubVLM.idx += 1
            return txt

    app = make_fake_app()
    app.word_path.set(docx)
    app.excel_path.set(xlsx)
    app.sheet_name.set("客戶資料")
    app.output_dir.set(tempfile.mkdtemp(prefix="review_out_"))
    app.enable_review.set(True)

    ctx = AppContext(app)
    ctx._build_reviewer_client = lambda: (StubVLM(), "stub-v", None)

    StubVLM.responses = [
        '{"passed": true, "score": 9, "issues": [], "suggestions": []}',
        '{"passed": false, "score": 4, "issues": ["bad"], "suggestions": []}',
        '{"passed": true, "score": 8, "issues": [], "suggestions": []}',
    ]
    StubVLM.idx = 0
    r = ctx.generate_reports()
    assert r["produced"] == 3
    assert r["reviewed"] == 3
    assert r["failed_count"] == 1
    print(f"  3 produced, 1 review failure → moved to {os.path.basename(r['failed_dir'])}")

    # Budget exhaustion
    app.budget = BudgetTracker(planner_limit=10, reviewer_limit=1)
    app.budget.reset()
    StubVLM.idx = 0
    r = ctx.generate_reports()
    assert r["produced"] == 3
    assert r["reviewed"] == 1
    assert r["review_budget_exhausted"]
    print(f"  budget=1: produced 3, reviewed 1, exhausted=True")


def test_orchestrator(docx, xlsx):
    section("orchestrator end-to-end")
    from app.agent.context import AppContext
    from app.agent.budget import BudgetTracker
    from app.agent.llm.base import Message, ToolCall
    from app.agent.orchestrator import AgentOrchestrator
    from app.agent.tools import build_default_registry

    app = make_fake_app()
    ctx = AppContext(app)
    reg = build_default_registry(ctx)

    class ScriptedLLM:
        def __init__(self, scripts): self.scripts = scripts; self.i = 0
        def chat(self, messages, model=None, tools=None):
            m = self.scripts[self.i]; self.i += 1; return m

    scripts = [
        Message(role="assistant", tool_calls=[ToolCall(name="get_current_settings", arguments={})]),
        Message(role="assistant", tool_calls=[ToolCall(name="list_excel_sheets", arguments={"path": xlsx})]),
        Message(role="assistant", text="DONE"),
    ]
    budget = BudgetTracker(10, 10)
    orch = AgentOrchestrator(ScriptedLLM(scripts), reg, model="t", budget=budget)
    orch.add_user_message("看設定 + 列 sheets")
    out = list(orch.step())
    assert [m.role for m in out] == ["assistant", "tool", "assistant", "tool", "assistant"]
    sheets_result = json.loads(out[3].text)
    assert sheets_result["sheets"] == ["客戶資料", "其他"]
    assert budget.planner_used == 3
    print(f"  multi-tool flow + budget tracking: ok")

    # Cancel
    orch2 = AgentOrchestrator(ScriptedLLM(scripts[:]), reg, model="t", budget=budget)
    orch2.add_user_message("test")
    orch2.cancel()
    out2 = list(orch2.step())
    assert out2[-1].text == "[已中止]"
    print(f"  cancel: ok")

    # Budget exhaustion stops planner
    class LoopLLM:
        def chat(self, messages, model=None, tools=None):
            return Message(role="assistant", tool_calls=[ToolCall(name="get_current_settings", arguments={})])
    budget2 = BudgetTracker(planner_limit=2, reviewer_limit=10)
    orch3 = AgentOrchestrator(LoopLLM(), reg, model="t", budget=budget2, max_iters=20)
    orch3.add_user_message("loop")
    msgs = list(orch3.step())
    assert any("[已達 planner 預算上限" in (m.text or "") for m in msgs)
    assert budget2.planner_used == 2
    print(f"  budget hard-stops planner: ok")


def test_llm_clients():
    section("LLM clients graceful degradation")
    from app.agent.llm.gemini import GeminiClient
    from app.agent.llm.ollama import OllamaClient

    g = GeminiClient(api_key="")
    assert not g.is_available()
    assert g.list_models() == []

    o = OllamaClient(endpoint="http://10.255.255.255:11434", timeout=1)
    assert not o.is_available()
    assert o.list_models() == []
    print(f"  no API key / unreachable endpoint: graceful")


def test_ui_structure():
    section("UI structure")
    import app.ui as ui
    cls = ui.AutoReportApp
    needed = [
        "_build_ui", "_build_settings_tab", "_build_mapping_tab",
        "_build_generate_tab", "_build_ai_tab", "_build_agent_tab",
        "_select_word", "_select_excel", "_select_output_dir",
        "_refresh_sheet_list",
        "_toggle_mapping", "_on_hotkey", "_undo_last_mapping",
        "_undo_all_mappings", "_insert_image_at_cursor", "_insert_image_grid_at_cursor",
        "_validate", "_toggle_generation", "_process_files",
        "_show_completion_summary", "_open_output_dir",
        "_show_provider_frame", "_on_provider_changed",
        "_test_llm", "_refresh_llm_models",
        "_reset_budget", "_refresh_budget_label", "_reset_rubric",
        "_refresh_generate_button_state",
        "_agent_send", "_agent_reset", "_agent_cancel",
        "_agent_export_chat", "_agent_render", "_agent_set_status",
        "_agent_box_block_keys",
        "_on_appearance_changed", "_on_file_drop", "_route_dropped_path",
        "_llm_ready",
    ]
    for m in needed:
        assert hasattr(cls, m), f"missing {m}"
    print(f"  all {len(needed)} UI methods present")
    assert callable(ui.open_folder)
    assert callable(ui._short_path)
    assert ui._short_path("") == ""
    assert ui._short_path("file.docx") == "file.docx"


# ============================================================
# Main
# ============================================================

def main():
    print("=" * 60)
    print("Auto-Report 預檢測試")
    print("=" * 60)

    test_imports()
    test_settings()
    test_filename()
    fixtures, docx, xlsx, img_dir, images = make_fixtures()
    print(f"\nfixtures: {fixtures}")
    test_generator(docx, xlsx)
    test_registry(docx, xlsx)
    test_template_edit(docx, images)
    test_folder_scan()
    test_reviewer()
    test_suggest_mapping(docx, xlsx, img_dir)
    test_review_flow(docx, xlsx)
    test_orchestrator(docx, xlsx)
    test_llm_clients()
    test_ui_structure()

    print("\n" + "=" * 60)
    print("全部通過")
    print("=" * 60)
    print("注意：")
    print("- 需 Windows 才能測 Word COM、真實 Tk 視窗、DnD、快捷鍵")
    print("- 需真實 Gemini/Ollama 才能測對話品質")


if __name__ == "__main__":
    main()
